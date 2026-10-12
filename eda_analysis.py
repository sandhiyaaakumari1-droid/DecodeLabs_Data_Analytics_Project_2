"""
DecodeLabs | Data Analytics | Project 2: Exploratory Data Analysis (EDA)

Input  : Cleaned_Dataset.xlsx (output of Project 1, sheet 'Cleaned_Data')
Outputs: charts/*.png, eda_stats.json, EDA_Statistics.xlsx

Usage (arguments are optional):
    python eda_analysis.py
    python eda_analysis.py <Cleaned_Dataset.xlsx> <output_dir>
With no arguments it reads Cleaned_Dataset.xlsx from the same folder as this script
and writes charts/, EDA_Statistics.xlsx and eda_stats.json there too.

Follows the IPO framework from the training kit:
    INPUT   -> sanitized data from Project 1
    PROCESS -> descriptive statistics, outliers (IQR vs Z-score), trends, segments, correlation
    OUTPUT  -> charts, statistics workbook, JSON for the report
"""
import sys
import os
import json

import numpy as np
import pandas as pd
from scipy import stats
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mtick
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "Cleaned_Dataset.xlsx")
OUT = sys.argv[2] if len(sys.argv) > 2 else BASE
CH = os.path.join(OUT, "charts")
os.makedirs(CH, exist_ok=True)

NAVY, TEAL, RED, GREY, LIGHT = "#1F3A5F", "#2A9D8F", "#C0392B", "#8A94A6", "#D9DEE7"
plt.rcParams.update({"font.family": "DejaVu Sans", "axes.spines.top": False,
                     "axes.spines.right": False, "axes.titleweight": "bold",
                     "axes.titlesize": 12, "figure.dpi": 150})

# =========================================================== INPUT
df = pd.read_excel(SRC, sheet_name="Cleaned_Data")
df["Year"] = df["Date"].dt.year
df["Month"] = df["Date"].dt.month
NUM = ["Quantity", "UnitPrice", "ItemsInCart", "TotalPrice"]
CATS = ["Product", "PaymentMethod", "ReferralSource", "CouponCode", "OrderStatus"]
tot = df["TotalPrice"].sum()
S = {}  # everything the report needs

# =========================================================== PROCESS
# ---- 1. Descriptive statistics (mean, median, count + five-number summary)
desc = {}
for c in NUM:
    x = df[c]
    desc[c] = {"count": int(x.count()), "mean": x.mean(), "median": x.median(), "std": x.std(),
               "min": x.min(), "q1": x.quantile(.25), "q3": x.quantile(.75), "max": x.max(),
               "skew": x.skew()}
S["desc"] = desc
S["rows"] = len(df)
S["date_min"] = df["Date"].min().strftime("%Y-%m-%d")
S["date_max"] = df["Date"].max().strftime("%Y-%m-%d")
S["total_value"] = tot
S["unique_customers"] = int(df["CustomerID"].nunique())
vc = df["CustomerID"].value_counts()
S["single_purchase_customers"] = int((vc == 1).sum())
S["single_purchase_pct"] = round((vc == 1).mean() * 100, 1)
S["shapiro_p_totalprice"] = float(stats.shapiro(df["TotalPrice"]).pvalue)
rng = np.random.default_rng(0)
boot = [np.median(rng.choice(df["TotalPrice"], len(df))) for _ in range(5000)]
S["median_ci95"] = [float(v) for v in np.percentile(boot, [2.5, 97.5])]
s_sorted = df["TotalPrice"].sort_values(ascending=False)
S["top10_share"] = round(s_sorted.head(int(len(s_sorted) * .1)).sum() / tot * 100, 1)
S["top20_share"] = round(s_sorted.head(int(len(s_sorted) * .2)).sum() / tot * 100, 1)
S["bottom50_share"] = round(s_sorted.tail(int(len(s_sorted) * .5)).sum() / tot * 100, 1)

# ---- 2. Outliers: IQR vs Z-score
out = {}
for c in NUM:
    q1, q3 = df[c].quantile([.25, .75]); iqr = q3 - q1
    lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
    z = np.abs(stats.zscore(df[c]))
    out[c] = {"lower_fence": lo, "upper_fence": hi, "iqr_count": int(((df[c] < lo) | (df[c] > hi)).sum()),
              "z_count": int((z > 3).sum()), "max_abs_z": float(z.max())}
S["outliers"] = out
hi = out["TotalPrice"]["upper_fence"]
o = df[df["TotalPrice"] > hi].sort_values("TotalPrice", ascending=False)
S["outlier_orders"] = o[["OrderID", "Product", "Quantity", "UnitPrice", "TotalPrice", "OrderStatus"]].to_dict("records")
S["outlier_value"] = float(o["TotalPrice"].sum())
S["outlier_value_pct"] = round(o["TotalPrice"].sum() / tot * 100, 2)
S["outlier_status"] = o["OrderStatus"].value_counts().to_dict()
S["outlier_qty_unique"] = sorted(o["Quantity"].unique().tolist())
S["outlier_price_range"] = [float(o["UnitPrice"].min()), float(o["UnitPrice"].max())]

# ---- 3. Trends
monthly = df.groupby(df["Date"].dt.to_period("M")).agg(orders=("OrderID", "count"), revenue=("TotalPrice", "sum"))
x = np.arange(len(monthly))
lr = stats.linregress(x, monthly["revenue"].values)
S["trend"] = {"slope_per_month": lr.slope, "p": lr.pvalue, "r2": lr.rvalue ** 2,
              "monthly_mean": float(monthly["revenue"].mean()), "monthly_cv": float(monthly["revenue"].std() / monthly["revenue"].mean())}
yr = df.groupby("Year").agg(orders=("OrderID", "count"), revenue=("TotalPrice", "sum"))
S["by_year"] = yr.to_dict("index")
h1 = df[df["Month"] <= 6].groupby("Year").agg(orders=("OrderID", "count"), revenue=("TotalPrice", "sum"))
h1["aov"] = h1["revenue"] / h1["orders"]
S["h1"] = h1.to_dict("index")
S["h1_yoy_rev_pct"] = {int(y): round((h1.loc[y, "revenue"] / h1.loc[y - 1, "revenue"] - 1) * 100, 1) for y in (2024, 2025)}
a, b = h1.loc[2023], h1.loc[2025]
S["h1_23_25"] = {"revenue_pct": round((b.revenue / a.revenue - 1) * 100, 1),
                 "orders_pct": round((b.orders / a.orders - 1) * 100, 1),
                 "aov_pct": round((b.aov / a.aov - 1) * 100, 1)}
fy = yr.copy()
S["fy_2024_vs_2023"] = {"revenue_pct": round((fy.loc[2024, "revenue"] / fy.loc[2023, "revenue"] - 1) * 100, 1),
                        "orders_pct": round((fy.loc[2024, "orders"] / fy.loc[2023, "orders"] - 1) * 100, 1)}
ym = df.groupby(["Year", "Month"])["TotalPrice"].sum().unstack(0)
S["june_rank_h1"] = {int(y): int(ym[y].dropna().loc[1:6].rank(ascending=False).loc[6]) for y in (2023, 2024, 2025)}
S["june_vs_h1_mean"] = {int(y): round(float(ym[y].dropna().loc[1:6].loc[6] / ym[y].dropna().loc[1:6].mean()), 2) for y in (2023, 2024, 2025)}
dow = df["Date"].dt.dayofweek
S["weekday_chi2_p"] = float(stats.chisquare(dow.value_counts()).pvalue)
S["weekday_kw_p"] = float(stats.kruskal(*[g["TotalPrice"].values for _, g in df.groupby(dow)]).pvalue)
S["moy_kw_p"] = float(stats.kruskal(*[g["TotalPrice"].values for _, g in df.groupby("Month")]).pvalue)

# ---- 4. Order status (revenue at risk)
st = df.groupby("OrderStatus").agg(orders=("OrderID", "count"), value=("TotalPrice", "sum"))
st["order_pct"] = st["orders"] / len(df) * 100
st["value_pct"] = st["value"] / tot * 100
S["status"] = st.round(2).to_dict("index")
risk = df[df["OrderStatus"].isin(["Cancelled", "Returned"])]
S["risk"] = {"orders": len(risk), "orders_pct": round(len(risk) / len(df) * 100, 1),
             "value": float(risk["TotalPrice"].sum()), "value_pct": round(risk["TotalPrice"].sum() / tot * 100, 1)}
infl = df[df["OrderStatus"].isin(["Pending", "Shipped"])]
S["in_flight"] = {"orders_pct": round(len(infl) / len(df) * 100, 1), "value": float(infl["TotalPrice"].sum()),
                  "value_pct": round(infl["TotalPrice"].sum() / tot * 100, 1)}
S["delivered"] = {"value": float(st.loc["Delivered", "value"]), "value_pct": round(st.loc["Delivered", "value_pct"], 1)}
S["risk_by_year"] = {int(y): round(g["OrderStatus"].isin(["Cancelled", "Returned"]).mean() * 100, 1) for y, g in df.groupby("Year")}
S["status_year_chi2_p"] = float(stats.chi2_contingency(pd.crosstab(df["Year"], df["OrderStatus"]))[1])
S["status_uniform_chi2_p"] = float(stats.chisquare(df["OrderStatus"].value_counts()).pvalue)

# ---- 5. Segments: do groups differ?
seg = {}
for c in CATS[:4]:
    g = df.groupby(c).agg(orders=("OrderID", "count"), revenue=("TotalPrice", "sum"),
                          mean=("TotalPrice", "mean"), median=("TotalPrice", "median"))
    seg[c] = {"table": g.round(2).to_dict("index"),
              "kruskal_p_value": float(stats.kruskal(*[x["TotalPrice"].values for _, x in df.groupby(c)]).pvalue),
              "chi2_volume_p": float(stats.chisquare(df[c].value_counts()).pvalue)}
S["segments"] = seg
used = df["CouponCode"] != "No Coupon"
S["coupon"] = {"used_orders": int(used.sum()), "used_pct": round(used.mean() * 100, 1),
               "median_used": float(df.loc[used, "TotalPrice"].median()), "median_none": float(df.loc[~used, "TotalPrice"].median()),
               "mean_used": float(df.loc[used, "TotalPrice"].mean()), "mean_none": float(df.loc[~used, "TotalPrice"].mean()),
               "mwu_p": float(stats.mannwhitneyu(df.loc[used, "TotalPrice"], df.loc[~used, "TotalPrice"]).pvalue)}
S["unitprice_by_product_mean"] = df.groupby("Product")["UnitPrice"].mean().round(0).to_dict()
S["unitprice_range"] = [float(df["UnitPrice"].min()), float(df["UnitPrice"].max())]

# ---- 6. Correlation
pear = df[NUM].corr()
S["pearson"] = pear.round(4).to_dict()
S["spearman"] = df[NUM].corr("spearman").round(4).to_dict()
S["items_lt_qty"] = int((df["ItemsInCart"] < df["Quantity"]).sum())
S["items_eq_qty"] = int((df["ItemsInCart"] == df["Quantity"]).sum())

# =========================================================== CHARTS
money = mtick.FuncFormatter(lambda v, _: f"${v:,.0f}")
kfmt = mtick.FuncFormatter(lambda v, _: f"${v/1000:,.0f}k")
mean_t, med_t = df["TotalPrice"].mean(), df["TotalPrice"].median()

# 1. distribution
fig, ax = plt.subplots(figsize=(8, 4.2))
ax.hist(df["TotalPrice"], bins=30, color=LIGHT, edgecolor="white")
ax.axvline(med_t, color=TEAL, lw=2.2); ax.axvline(mean_t, color=RED, lw=2.2, ls="--")
ymax = ax.get_ylim()[1]
ax.text(med_t - 60, ymax * .93, f"Median\n${med_t:,.0f}", color=TEAL, ha="right", va="top", fontweight="bold", fontsize=9)
ax.text(mean_t + 60, ymax * .93, f"Mean\n${mean_t:,.0f}", color=RED, ha="left", va="top", fontweight="bold", fontsize=9)
ax.text(.98, .55, f"Skewness = {desc['TotalPrice']['skew']:.2f}\n(right-skewed)", transform=ax.transAxes, ha="right", fontsize=9, color=NAVY)
ax.set_title(f"A typical order is \\${med_t:,.0f}; a long tail pulls the average up to \\${mean_t:,.0f}", loc="left")  # \\$ escapes matplotlib math mode
ax.set_xlabel("Order value (TotalPrice)"); ax.set_ylabel("Number of orders"); ax.xaxis.set_major_formatter(money)
fig.tight_layout(); fig.savefig(f"{CH}/01_order_value_distribution.png"); plt.close(fig)

# 2. outliers boxplot
fig, ax = plt.subplots(figsize=(8, 3.3))
bp = ax.boxplot(df["TotalPrice"], vert=False, widths=.5, patch_artist=True, whis=1.5,
                boxprops=dict(facecolor=LIGHT, color=NAVY), medianprops=dict(color=TEAL, lw=2.5),
                whiskerprops=dict(color=NAVY), capprops=dict(color=NAVY),
                flierprops=dict(marker="o", markerfacecolor=RED, markeredgecolor=RED, markersize=6))
ax.axvline(hi, color=RED, ls=":", lw=1.4)
ax.text(hi, 1.42, f"IQR fence ${hi:,.0f}", color=RED, ha="center", fontsize=8.5)
ax.set_yticks([]); ax.xaxis.set_major_formatter(money); ax.set_xlabel("Order value (TotalPrice)")
ax.text(.01, .06, f"IQR rule flags {out['TotalPrice']['iqr_count']} orders  |  Z-score rule (|z| > 3) flags {out['TotalPrice']['z_count']}",
        transform=ax.transAxes, fontsize=9, color=NAVY, fontweight="bold")
ax.set_title("The IQR rule flags 8 unusually large orders that the Z-score rule misses", loc="left")
ax.spines["left"].set_visible(False)
fig.tight_layout(); fig.savefig(f"{CH}/02_outliers_boxplot.png"); plt.close(fig)

# 3. monthly trend
fig, ax = plt.subplots(figsize=(8.6, 4.2))
idx = monthly.index.to_timestamp()
ax.plot(idx, monthly["revenue"], color=GREY, marker="o", ms=3.5, lw=1, label="Monthly revenue")
ax.plot(idx, monthly["revenue"].rolling(3).mean(), color=NAVY, lw=2.6, label="3-month average")
ax.plot(idx, lr.intercept + lr.slope * x, color=RED, ls="--", lw=1.6,
        label=f"Linear trend (about -${abs(lr.slope):,.0f}/month, p = {lr.pvalue:.2f}, not significant)")
for i, d in enumerate(idx):
    if d.month == 6:
        ax.scatter([d], [monthly["revenue"].iloc[i]], color=TEAL, s=55, zorder=5)
        ax.annotate("Jun", (d, monthly["revenue"].iloc[i]), textcoords="offset points", xytext=(0, 7), ha="center", fontsize=8, color=TEAL, fontweight="bold")
ax.yaxis.set_major_formatter(kfmt); ax.set_ylim(0, 75000); ax.legend(frameon=False, fontsize=8, loc="lower left")
ax.set_title("Monthly revenue is volatile but drifting down", loc="left"); ax.set_ylabel("Revenue")
fig.tight_layout(); fig.savefig(f"{CH}/03_monthly_revenue_trend.png"); plt.close(fig)

# 4. H1 year over year
fig, ax = plt.subplots(figsize=(6.4, 3.8))
yrs = [2023, 2024, 2025]; vals = [h1.loc[y, "revenue"] for y in yrs]
bars = ax.bar([f"H1 {y}" for y in yrs], vals, color=[LIGHT, LIGHT, NAVY], width=.55)
for bar, v in zip(bars, vals):
    ax.text(bar.get_x() + bar.get_width() / 2, v + 4000, f"${v/1000:,.1f}k", ha="center", fontweight="bold", fontsize=10)
for i, y in enumerate([2024, 2025], start=1):
    ax.text(i, vals[i] / 2, f"{S['h1_yoy_rev_pct'][y]:+.1f}%", ha="center", color=("white" if i == 2 else RED), fontweight="bold", fontsize=12)
ax.yaxis.set_major_formatter(kfmt); ax.set_ylim(0, 340000)
ax.set_title("First-half revenue fell about 10% in each of the last two years", loc="left", fontsize=11)
fig.tight_layout(); fig.savefig(f"{CH}/04_h1_year_over_year.png"); plt.close(fig)

# 5. order status
order = st.sort_values("value_pct")
cols = [RED if s in ("Cancelled", "Returned") else (TEAL if s == "Delivered" else GREY) for s in order.index]
fig, ax = plt.subplots(figsize=(8, 3.9))
bars = ax.barh(order.index, order["value_pct"], color=cols, height=.6)
for bar, (s, r) in zip(bars, order.iterrows()):
    ax.text(r["value_pct"] + .3, bar.get_y() + bar.get_height() / 2, f"{r['value_pct']:.1f}%  (${r['value']/1000:,.0f}k)", va="center", fontsize=9)
ax.set_xlim(0, 28); ax.xaxis.set_major_formatter(mtick.PercentFormatter(decimals=0)); ax.set_xlabel("Share of total order value")
ax.set_title(f"{S['risk']['value_pct']:.0f}% of order value was cancelled or returned; only {S['delivered']['value_pct']:.0f}% is delivered", loc="left", fontsize=11)
fig.tight_layout(); fig.savefig(f"{CH}/05_order_status.png"); plt.close(fig)

# 6. segments (2x2)
fig, axes = plt.subplots(2, 2, figsize=(9, 6.2), sharey=True)
for ax, c in zip(axes.ravel(), ["Product", "PaymentMethod", "ReferralSource", "CouponCode"]):
    t = pd.DataFrame(seg[c]["table"]).T.sort_values("median", ascending=False)
    ax.bar(t.index, t["median"], color=LIGHT, edgecolor=GREY)
    ax.axhline(med_t, color=TEAL, ls="--", lw=1.5)
    ax.set_title(f"{c}   (p = {seg[c]['kruskal_p_value']:.2f})", loc="left", fontsize=10)
    ax.tick_params(axis="x", rotation=35, labelsize=8); ax.yaxis.set_major_formatter(money)
    for lab in ax.get_xticklabels(): lab.set_ha("right")
axes[0, 0].set_ylabel("Median order value"); axes[1, 0].set_ylabel("Median order value")
fig.suptitle("No segment stands out: order value is statistically the same across groups", x=.01, ha="left", fontweight="bold", fontsize=12)
fig.text(.01, .005, f"Dashed line = overall median (${med_t:,.0f}). p = Kruskal-Wallis test; p > 0.05 means no detectable difference.", fontsize=8, color=GREY)
fig.tight_layout(rect=[0, .02, 1, .95]); fig.savefig(f"{CH}/06_segment_comparison.png"); plt.close(fig)

# 7. correlation heatmap
fig, ax = plt.subplots(figsize=(6.4, 5.2))
im = ax.imshow(pear.values, cmap="coolwarm", vmin=-1, vmax=1)
ax.set_xticks(range(4)); ax.set_yticks(range(4)); ax.set_xticklabels(NUM, rotation=30, ha="right"); ax.set_yticklabels(NUM)
for i in range(4):
    for j in range(4):
        v = pear.values[i, j]; ax.text(j, i, f"{v:.2f}", ha="center", va="center", color="white" if abs(v) > .55 else "black", fontweight="bold")
ax.spines[:].set_visible(False); fig.colorbar(im, fraction=.046, pad=.04)
ax.set_title("Price and quantity drive order value, but not each other", loc="left", fontsize=11)
fig.tight_layout(); fig.savefig(f"{CH}/07_correlation_heatmap.png"); plt.close(fig)

# =========================================================== OUTPUT: workbook (live formulas)
FONT = "Arial"
HF = Font(name=FONT, bold=True, color="FFFFFF", size=10)
HFILL = PatternFill("solid", start_color="1F3A5F")
BF = Font(name=FONT, size=10)
BB = Font(name=FONT, size=10, bold=True)
BLUE = Font(name=FONT, size=10, color="0000FF")
TF = Font(name=FONT, size=13, bold=True, color="1F3A5F")
NF = Font(name=FONT, size=9, italic=True, color="595959")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def hdr(ws, row, labels, col=1):
    for j, l in enumerate(labels, start=col):
        c = ws.cell(row, j, l); c.font, c.fill, c.border = HF, HFILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def put(ws, row, col, val, fmt=None, font=BF, align=None):
    c = ws.cell(row, col, val); c.font = font; c.border = BOX
    if fmt: c.number_format = fmt
    if align: c.alignment = Alignment(horizontal=align)
    return c

wb = Workbook()
# ---- Data
wd = wb.active; wd.title = "Data"
cols = list(pd.read_excel(SRC, sheet_name="Cleaned_Data").columns)
raw = pd.read_excel(SRC, sheet_name="Cleaned_Data")
wd.append(cols)
for rec in raw.itertuples(index=False):
    wd.append([v.to_pydatetime() if isinstance(v, pd.Timestamp) else (int(v) if isinstance(v, (np.integer,)) else (float(v) if isinstance(v, np.floating) else v)) for v in rec])
n = len(raw); last = n + 1
L = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
R = lambda c: f"Data!${L[c]}$2:${L[c]}${last}"
hdr(wd, 1, cols)
for r in range(2, last + 1):
    for j, c in enumerate(cols, start=1):
        cell = wd.cell(r, j); cell.font = BF
        if c == "Date": cell.number_format = "yyyy-mm-dd"
        elif c in ("UnitPrice", "TotalPrice"): cell.number_format = "#,##0.00"
for j, c in enumerate(cols, start=1): wd.column_dimensions[get_column_letter(j)].width = max(12, len(c) + 3)
wd.freeze_panes = "A2"; wd.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

# ---- Descriptive_Stats
ws = wb.create_sheet("Descriptive_Stats")
ws["A1"] = "Descriptive Statistics, Five-Number Summary and Outlier Rules"; ws["A1"].font = TF
ws["A2"] = "All values are live formulas on the Data sheet. IQR rule: outside Q1 - 1.5 x IQR or Q3 + 1.5 x IQR. Z-score rule: |z| > 3."; ws["A2"].font = NF
hdr(ws, 4, ["Statistic"] + NUM)
metrics = [
    ("Count", "=COUNT({r})", "#,##0"), ("Mean", "=AVERAGE({r})", "#,##0.00"), ("Median", "=MEDIAN({r})", "#,##0.00"),
    ("Mean minus Median", "={c}6-{c}7", "#,##0.00"),
    ("Std deviation", "=STDEV({r})", "#,##0.00"), ("Minimum", "=MIN({r})", "#,##0.00"),
    ("Q1 (25th percentile)", "=QUARTILE({r},1)", "#,##0.00"), ("Q3 (75th percentile)", "=QUARTILE({r},3)", "#,##0.00"),
    ("Maximum", "=MAX({r})", "#,##0.00"), ("IQR (Q3 - Q1)", "={c}12-{c}11", "#,##0.00"),
    ("Skewness", "=SKEW({r})", "0.00"),
    ("IQR lower fence", "={c}11-1.5*{c}14", "#,##0.00"), ("IQR upper fence", "={c}12+1.5*{c}14", "#,##0.00"),
    ("Outliers by IQR rule", '=COUNTIF({r},"<"&{c}16)+COUNTIF({r},">"&{c}17)', "0"),
    ("Outliers by Z-score rule (|z|>3)", "=SUMPRODUCT(--(ABS(({r}-{c}6)/{c}9)>3))", "0"),
]
for i, (name, f, fmt) in enumerate(metrics, start=5):
    put(ws, i, 1, name, font=BB)
    for j, c in enumerate(NUM, start=2):
        put(ws, i, j, f.format(r=R(c), c=get_column_letter(j)), fmt, align="right")
ws.column_dimensions["A"].width = 34
for j in range(2, 6): ws.column_dimensions[get_column_letter(j)].width = 15
ws["A21"] = "Reading guide: skewness near 0 = symmetric (mean is a fair summary); above ~0.5 = right-skewed (use the median)."; ws["A21"].font = NF
ws.freeze_panes = "B5"

# ---- By_Category
wc = wb.create_sheet("By_Category")
wc["A1"] = "Orders and Order Value by Segment"; wc["A1"].font = TF
wc["A2"] = "Median order value uses array formulas (MEDIAN(IF(...))). Counts, sums and averages are live formulas on the Data sheet."; wc["A2"].font = NF
row = 4
for c in CATS:
    wc.cell(row, 1, c).font = Font(name=FONT, size=11, bold=True, color="1F3A5F"); row += 1
    hdr(wc, row, ["Group", "Orders", "Share of orders", "Order value", "Share of value", "Mean order value", "Median order value"]); row += 1
    groups = sorted(df[c].unique()); first = row
    for g in groups:
        put(wc, row, 1, g, font=BB)
        put(wc, row, 2, f'=COUNTIF({R(c)},A{row})', "#,##0", align="right")
        put(wc, row, 3, f"=B{row}/COUNTA({R('OrderID')})", "0.0%", align="right")
        put(wc, row, 4, f'=SUMIF({R(c)},A{row},{R("TotalPrice")})', "$#,##0", align="right")
        put(wc, row, 5, f"=D{row}/SUM({R('TotalPrice')})", "0.0%", align="right")
        put(wc, row, 6, f'=AVERAGEIF({R(c)},A{row},{R("TotalPrice")})', "$#,##0.00", align="right")
        cell = wc.cell(row, 7); cell.value = ArrayFormula(f"G{row}", f'=MEDIAN(IF({R(c)}=A{row},{R("TotalPrice")}))')
        cell.number_format = "$#,##0.00"; cell.font = BF; cell.border = BOX; cell.alignment = Alignment(horizontal="right")
        row += 1
    put(wc, row, 1, "Total", font=BB)
    put(wc, row, 2, f"=SUM(B{first}:B{row-1})", "#,##0", BB, "right"); put(wc, row, 3, f"=SUM(C{first}:C{row-1})", "0.0%", BB, "right")
    put(wc, row, 4, f"=SUM(D{first}:D{row-1})", "$#,##0", BB, "right"); put(wc, row, 5, f"=SUM(E{first}:E{row-1})", "0.0%", BB, "right")
    row += 2
wc.column_dimensions["A"].width = 22
for j in range(2, 8): wc.column_dimensions[get_column_letter(j)].width = 17
# cancelled + returned summary
wc.cell(row, 1, "Revenue at risk").font = Font(name=FONT, size=11, bold=True, color="1F3A5F"); row += 1
put(wc, row, 1, "Cancelled + Returned value", font=BB)
put(wc, row, 4, f'=SUMIF({R("OrderStatus")},"Cancelled",{R("TotalPrice")})+SUMIF({R("OrderStatus")},"Returned",{R("TotalPrice")})', "$#,##0", BB, "right")
put(wc, row, 5, f"=D{row}/SUM({R('TotalPrice')})", "0.0%", BB, "right")
put(wc, row, 2, f'=COUNTIF({R("OrderStatus")},"Cancelled")+COUNTIF({R("OrderStatus")},"Returned")', "#,##0", BB, "right")
put(wc, row, 3, f"=B{row}/COUNTA({R('OrderID')})", "0.0%", BB, "right")

# ---- Monthly_Trend
wm = wb.create_sheet("Monthly_Trend")
wm["A1"] = "Monthly Orders and Revenue, with First-Half Year-over-Year Comparison"; wm["A1"].font = TF
hdr(wm, 3, ["Month start", "Orders", "Revenue", "3-month avg revenue"])
months = pd.period_range(df["Date"].min(), df["Date"].max(), freq="M")
for i, p in enumerate(months, start=4):
    put(wm, i, 1, p.to_timestamp().to_pydatetime(), "yyyy-mm", align="center")
    put(wm, i, 2, f'=COUNTIFS({R("Date")},">="&A{i},{R("Date")},"<"&EDATE(A{i},1))', "0", align="right")
    put(wm, i, 3, f'=SUMIFS({R("TotalPrice")},{R("Date")},">="&A{i},{R("Date")},"<"&EDATE(A{i},1))', "$#,##0", align="right")
    put(wm, i, 4, f"=AVERAGE(C{i-2}:C{i})" if i >= 6 else None, "$#,##0", align="right")
lm = 3 + len(months)
put(wm, lm + 2, 1, "Trend slope ($/month)", font=BB); put(wm, lm + 2, 3, f"=SLOPE(C4:C{lm},ROW(C4:C{lm})-3)", "#,##0.0", align="right")
put(wm, lm + 3, 1, "Trend R-squared", font=BB); put(wm, lm + 3, 3, f"=RSQ(C4:C{lm},ROW(C4:C{lm})-3)", "0.000", align="right")
put(wm, lm + 4, 1, "Monthly revenue mean", font=BB); put(wm, lm + 4, 3, f"=AVERAGE(C4:C{lm})", "$#,##0", align="right")
wm.cell(lm + 5, 1, "A low R-squared means month-to-month noise dominates the trend; see Stat_Tests for the p-value.").font = NF
# H1 table
hdr(wm, 3, ["First half (Jan-Jun)", "Orders", "Revenue", "Mean order value", "Revenue vs prior year"], col=7)
for k, y in enumerate((2023, 2024, 2025)):
    r = 4 + k
    put(wm, r, 7, y, "0", BB, "center")
    crit = f'{R("Date")},">="&DATE(G{r},1,1),{R("Date")},"<"&DATE(G{r},7,1)'
    put(wm, r, 8, f"=COUNTIFS({crit})", "0", align="right")
    put(wm, r, 9, f'=SUMIFS({R("TotalPrice")},{crit})', "$#,##0", align="right")
    put(wm, r, 10, f"=I{r}/H{r}", "$#,##0.00", align="right")
    put(wm, r, 11, f"=I{r}/I{r-1}-1" if k else None, "0.0%", align="right")
put(wm, 8, 7, "2023 to 2025", font=BB, align="center")
put(wm, 8, 8, "=H6/H4-1", "0.0%", align="right"); put(wm, 8, 9, "=I6/I4-1", "0.0%", align="right"); put(wm, 8, 10, "=J6/J4-1", "0.0%", align="right")
wm.cell(9, 7, "Revenue change = order-volume change x mean-order-value change.").font = NF
wm.column_dimensions["A"].width = 24
for col, w in zip("BCD", [10, 14, 20]): wm.column_dimensions[col].width = w
for col, w in zip("GHIJK", [20, 10, 14, 18, 22]): wm.column_dimensions[col].width = w
wm.freeze_panes = "A4"

# ---- Correlation
wr = wb.create_sheet("Correlation")
wr["A1"] = "Pearson Correlation Matrix (r)"; wr["A1"].font = TF
wr["A2"] = "TotalPrice = Quantity x UnitPrice by construction, so its correlation with those two columns is expected, not a discovery."; wr["A2"].font = NF
hdr(wr, 4, [""] + NUM)
for i, a_ in enumerate(NUM, start=5):
    put(wr, i, 1, a_, font=BB)
    for j, b_ in enumerate(NUM, start=2):
        put(wr, i, j, f"=CORREL({R(a_)},{R(b_)})", "0.00", align="center")
wr.column_dimensions["A"].width = 16
for j in range(2, 6): wr.column_dimensions[get_column_letter(j)].width = 14
from openpyxl.formatting.rule import ColorScaleRule
wr.conditional_formatting.add("B5:E8", ColorScaleRule(start_type="num", start_value=-1, start_color="2C5AA0",
                                                    mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=1, end_color="C0392B"))

# ---- Outliers
wo = wb.create_sheet("Outliers")
wo["A1"] = "Orders Flagged by the IQR Rule (TotalPrice)"; wo["A1"].font = TF
wo["A2"] = "Order list extracted by eda_analysis.py (blue = static). Upper fence and flag are live formulas."; wo["A2"].font = NF
put(wo, 3, 1, "Upper fence", font=BB); put(wo, 3, 2, "=Descriptive_Stats!E17", "$#,##0.00", BB, "right")
hdr(wo, 5, ["OrderID", "Product", "Quantity", "UnitPrice", "TotalPrice", "OrderStatus", "Above fence?"])
for i, rec in enumerate(S["outlier_orders"], start=6):
    for j, k in enumerate(["OrderID", "Product", "Quantity", "UnitPrice", "TotalPrice", "OrderStatus"], start=1):
        put(wo, i, j, rec[k], "$#,##0.00" if k in ("UnitPrice", "TotalPrice") else None, BLUE)
    put(wo, i, 7, f'=IF(E{i}>$B$3,"Yes","No")', align="center")
for col, w in zip("ABCDEFG", [14, 12, 10, 12, 13, 14, 14]): wo.column_dimensions[col].width = w

# ---- Stat_Tests (static results from scipy)
wt = wb.create_sheet("Stat_Tests")
wt["A1"] = "Hypothesis Tests (computed in Python / SciPy; static values in blue)"; wt["A1"].font = TF
wt["A2"] = "p < 0.05 would indicate a statistically detectable difference or trend. Every test below except the skew check is NOT significant."; wt["A2"].font = NF
hdr(wt, 4, ["Question", "Test", "p-value", "Conclusion"])
tests = [
    ("Does order value differ by Product?", "Kruskal-Wallis", S["segments"]["Product"]["kruskal_p_value"]),
    ("Does order value differ by Payment Method?", "Kruskal-Wallis", S["segments"]["PaymentMethod"]["kruskal_p_value"]),
    ("Does order value differ by Referral Source?", "Kruskal-Wallis", S["segments"]["ReferralSource"]["kruskal_p_value"]),
    ("Does order value differ by Coupon Code?", "Kruskal-Wallis", S["segments"]["CouponCode"]["kruskal_p_value"]),
    ("Do coupon orders have a different value than no-coupon orders?", "Mann-Whitney U", S["coupon"]["mwu_p"]),
    ("Does order value differ by weekday?", "Kruskal-Wallis", S["weekday_kw_p"]),
    ("Does order value differ by calendar month?", "Kruskal-Wallis", S["moy_kw_p"]),
    ("Are orders unevenly spread across weekdays?", "Chi-square (uniform)", S["weekday_chi2_p"]),
    ("Is the order-status mix uneven?", "Chi-square (uniform)", S["status_uniform_chi2_p"]),
    ("Did the status mix change by year?", "Chi-square (independence)", S["status_year_chi2_p"]),
    ("Is there a linear trend in monthly revenue?", "Linear regression", S["trend"]["p"]),
    ("Is TotalPrice normally distributed?", "Shapiro-Wilk", S["shapiro_p_totalprice"]),
]
for i, (q, t, p) in enumerate(tests, start=5):
    put(wt, i, 1, q); put(wt, i, 2, t); put(wt, i, 3, p, "0.000", BLUE, "center")
    put(wt, i, 4, f'=IF(C{i}<0.05,"Significant","Not significant")', align="center")
wt.column_dimensions["A"].width = 58; wt.column_dimensions["B"].width = 26; wt.column_dimensions["C"].width = 11; wt.column_dimensions["D"].width = 16

for w_ in (ws, wc, wm, wr, wo, wt): w_.sheet_view.showGridLines = False
wb.save(os.path.join(OUT, "EDA_Statistics.xlsx"))

with open(os.path.join(OUT, "eda_stats.json"), "w") as f:
    json.dump(S, f, indent=2, default=lambda o: int(o) if isinstance(o, np.integer) else (float(o) if isinstance(o, np.floating) else str(o)))
print("done")